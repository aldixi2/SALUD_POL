# -*- coding: utf-8 -*-
"""
generar_data.py
---------------
Lee el Excel ATENCIONES_SALUD_POL.xlsx directamente desde la carpeta local
y genera el data.json que alimenta el dashboard (index.html).

Uso (doble clic o desde la terminal):
    python generar_data.py

Requisito (una sola vez):
    pip install openpyxl

Hojas esperadas en el Excel: "CONSOLIDADO" y "Transferencias"
(las mismas columnas de siempre).
"""

import json
import re
import sys
import unicodedata
from datetime import datetime, date
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Falta openpyxl. Instalalo con:  pip install openpyxl")

# ======================================================================
# CONFIGURACION: cambia solo esto si mueves las carpetas
# ======================================================================
RUTA_BASE = Path(r"D:\ANALISIS DE DATOS\SALUD POL")
NOMBRE_EXCEL = "ATENCIONES_SALUD_POL.xlsx"

# Carpeta donde se guardara data.json (la del dashboard / repositorio).
# Por defecto: la misma carpeta donde esta este script.
CARPETA_SALIDA = Path(__file__).resolve().parent
# ======================================================================


def norm(s):
    """Normaliza encabezados: minusculas, sin tildes, sin guiones ni simbolos."""
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("-", "")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def to_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def r2(x):
    return round(x + 0.0, 2)


def fmt_fecha(v):
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return "" if v is None else str(v).strip()


def anio_de_fecha(v):
    return v.year if isinstance(v, (datetime, date)) else None


def limpiar_dni(v):
    """77837859.0 -> '77837859'; completa a 8 digitos si perdio ceros a la izquierda."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    if s.isdigit() and len(s) < 8:
        s = s.zfill(8)
    return s


def texto(v):
    return "" if v is None else str(v).strip()


MESES = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9,
    "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
}


def periodo_respaldo(mes, anio):
    """Si 'MES DE PRODUCCION' viene vacio, lo arma con las columnas MES y AÑO -> 'MM-YYYY'."""
    n = MESES.get(texto(mes).upper())
    return f"{n:02d}-{anio}" if n and anio else ""


def resolver_columnas(headers, mapa, fallbacks):
    hn = [norm(h) for h in headers]
    idx = {}
    for clave, objetivo in mapa.items():
        i = hn.index(objetivo) if objetivo in hn else -1
        if i == -1 and clave in fallbacks:
            for j, h in enumerate(hn):
                if all(t in h for t in fallbacks[clave]):
                    i = j
                    break
        idx[clave] = i
        if i == -1:
            print(f'  [!] No se encontro la columna "{clave}" (se esperaba "{objetivo}")')
        else:
            print(f'  [ok] {clave:14s} -> "{headers[i]}"')
    return idx


def buscar_hoja(wb, fragmento, defecto=None):
    for n in wb.sheetnames:
        if fragmento in norm(n):
            return wb[n]
    return defecto


def main():
    ruta_excel = RUTA_BASE / NOMBRE_EXCEL

    # Permite pasar otra ruta por argumento:  python generar_data.py "C:\otra\ruta.xlsx"
    if len(sys.argv) > 1:
        ruta_excel = Path(sys.argv[1])

    if not ruta_excel.exists():
        sys.exit(f"No encuentro el Excel en:\n  {ruta_excel}\n"
                 f"Revisa RUTA_BASE / NOMBRE_EXCEL al inicio de este archivo.")

    print(f"Leyendo: {ruta_excel}")
    wb = load_workbook(ruta_excel, read_only=True, data_only=True)
    print("Hojas encontradas:", wb.sheetnames)

    ws_c = buscar_hoja(wb, "consolidado", wb[wb.sheetnames[0]])
    ws_t = buscar_hoja(wb, "transferencia")

    # ---------------- CONSOLIDADO (atenciones) ----------------
    filas = list(ws_c.iter_rows(values_only=True))
    headers = [texto(h) for h in filas[0]]

    mapa = {
        "mesProduccion": "mes de produccion",
        "documento": "documento de autorizacion",
        "dni": "dni",
        "paciente": "nombre y apellido paciente",
        "montoProced": "monto procedimientos",
        "montoMedic": "monto medicamentos",
        "montoConforme": "monto conforme total",
        "fechaFactura": "fecha factura",
        "nroFactura": "nro factura",
        "montoFactura": "monto factura",
        "mes": "mes",
        "anio": "ano",
    }
    fb = {
        "mesProduccion": ["produ"], "documento": ["documento"],
        "montoProced": ["procedi"], "montoMedic": ["medica"],
        "montoConforme": ["conforme"], "fechaFactura": ["fecha", "factura"],
        "nroFactura": ["nro", "factura"], "montoFactura": ["monto", "factura"],
        "paciente": ["paciente"],
    }
    print("--- Columnas (atenciones) ---")
    ix = resolver_columnas(headers, mapa, fb)

    def get(row, clave):
        i = ix[clave]
        return row[i] if 0 <= i < len(row) else None

    atenciones = []
    sin_mes_prod = []
    for row in filas[1:]:
        if not row:
            continue
        dni_raw = get(row, "dni")
        if dni_raw is None or str(dni_raw).strip() == "":
            continue  # fila de totales o vacia

        f_raw = get(row, "fechaFactura")
        anio_raw = get(row, "anio")
        try:
            anio = int(float(anio_raw))
        except (TypeError, ValueError):
            anio = anio_de_fecha(f_raw)

        mes_prod = texto(get(row, "mesProduccion"))
        if not re.search(r"\d{1,2}[^\d]+\d{4}", mes_prod):
            mes_prod = periodo_respaldo(get(row, "mes"), anio) or mes_prod
            sin_mes_prod.append(texto(get(row, "documento")))

        atenciones.append({
            "mesProduccion": mes_prod,
            "documento": texto(get(row, "documento")),
            "dni": limpiar_dni(dni_raw),
            "paciente": texto(get(row, "paciente")),
            "montoProced": to_num(get(row, "montoProced")),
            "montoMedic": to_num(get(row, "montoMedic")),
            "montoConforme": to_num(get(row, "montoConforme")),
            "fechaFactura": fmt_fecha(f_raw),
            "nroFactura": texto(get(row, "nroFactura")),
            "montoFactura": to_num(get(row, "montoFactura")),
            "mes": get(row, "mes"),
            "anio": anio,
        })
    print(f"Atenciones validas: {len(atenciones)}")
    if sin_mes_prod:
        print(f"[i] {len(sin_mes_prod)} filas sin 'MES DE PRODUCCION': se uso MES + AÑO como periodo.")

    # ---------------- Transferencias ----------------
    transferencias = []
    if ws_t is None:
        print("[!] No se encontro la hoja de Transferencias.")
    else:
        filas_t = list(ws_t.iter_rows(values_only=True))
        headers_t = [texto(h) for h in filas_t[0]]
        print("--- Columnas (transferencias) ---")
        it = resolver_columnas(
            headers_t, {"n": "n", "monto": "depositos", "orden": "numero de orden"}, {}
        )
        for row in filas_t[1:]:
            if not row:
                continue
            n_val = row[it["n"]] if it["n"] > -1 else None
            try:
                n = int(float(n_val))
            except (TypeError, ValueError):
                continue  # fila de total u otra
            monto = to_num(row[it["monto"]]) if it["monto"] > -1 else 0.0
            orden = row[it["orden"]] if it["orden"] > -1 else ""
            if isinstance(orden, float) and orden.is_integer():
                orden = int(orden)
            transferencias.append({"n": n, "monto": r2(monto), "orden": texto(orden)})
        print(f"Transferencias validas: {len(transferencias)}")

    wb.close()

    # ---------------- Totales, mensual, anual ----------------
    transf_total = r2(sum(t["monto"] for t in transferencias))
    tot_factura = r2(sum(a["montoFactura"] for a in atenciones))
    tot_proced = r2(sum(a["montoProced"] for a in atenciones))
    tot_medic = r2(sum(a["montoMedic"] for a in atenciones))
    pacientes = len({a["dni"] for a in atenciones})
    pendiente = r2(tot_factura - transf_total)
    cobertura = r2(transf_total / tot_factura * 100) if tot_factura else 0

    mensual_map = {}
    for a in atenciones:
        m = re.search(r"(\d{1,2})[^\d]+(\d{4})", a["mesProduccion"])
        if not m:
            continue
        key = f"{m.group(2)}-{m.group(1).zfill(2)}"
        d = mensual_map.setdefault(key, {"periodo": key, "factura": 0, "proced": 0, "medic": 0, "n": 0})
        d["factura"] += a["montoFactura"]
        d["proced"] += a["montoProced"]
        d["medic"] += a["montoMedic"]
        d["n"] += 1
    mensual = [
        {"periodo": d["periodo"], "factura": r2(d["factura"]), "proced": r2(d["proced"]),
         "medic": r2(d["medic"]), "n": d["n"]}
        for d in sorted(mensual_map.values(), key=lambda x: x["periodo"])
    ]

    anual_map = {}
    for a in atenciones:
        y = a["anio"]
        if not y:
            continue
        d = anual_map.setdefault(y, {"anio": y, "factura": 0, "proced": 0, "medic": 0, "n": 0, "pac": set()})
        d["factura"] += a["montoFactura"]
        d["proced"] += a["montoProced"]
        d["medic"] += a["montoMedic"]
        d["n"] += 1
        d["pac"].add(a["dni"])
    anual = [
        {"anio": d["anio"], "factura": r2(d["factura"]), "proced": r2(d["proced"]),
         "medic": r2(d["medic"]), "n": d["n"], "pacientes": len(d["pac"])}
        for d in sorted(anual_map.values(), key=lambda x: x["anio"])
    ]

    data = {
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "periodoMin": mensual[0]["periodo"] if mensual else "",
        "periodoMax": mensual[-1]["periodo"] if mensual else "",
        "totales": {
            "factura": tot_factura, "proced": tot_proced, "medic": tot_medic,
            "atenciones": len(atenciones), "pacientes": pacientes,
            "transferido": transf_total, "pendiente": pendiente, "cobertura": cobertura,
        },
        "mensual": mensual,
        "anual": anual,
        "transferencias": transferencias,
        "atenciones": atenciones,
    }

    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)
    salida = CARPETA_SALIDA / "data.json"
    with open(salida, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))

    print("\n=========== RESUMEN ===========")
    print(f"Periodo        : {data['periodoMin']} a {data['periodoMax']}")
    print(f"Atenciones     : {len(atenciones):,}")
    print(f"Pacientes      : {pacientes:,}")
    print(f"Facturado      : S/ {tot_factura:,.2f}")
    print(f"Transferido    : S/ {transf_total:,.2f}")
    print(f"Pendiente      : S/ {pendiente:,.2f}")
    print(f"Cobertura      : {cobertura:.2f}%")
    print(f"\ndata.json guardado en:\n  {salida}")


if __name__ == "__main__":
    main()
